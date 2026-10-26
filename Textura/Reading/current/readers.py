"""Bounded, read-only text projections with honest source locators."""
from __future__ import annotations

import csv
import importlib.util
import re
from html.parser import HTMLParser
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_ZIP_BYTES = 128 * 1024 * 1024
MAX_MEMBER_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBERS = 10000
MAX_OUTPUT_CHARS = 2_000_000
MAX_BLOCKS = 50000
MAX_PAGES = 2000
MAX_SHEET_ROWS = 20000
MAX_SHEET_COLUMNS = 2000
MAX_CELLS = 200000
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class ReaderLimit(ValueError):
    pass


def _available(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError, AttributeError):
        return False


def capabilities() -> dict:
    """Advertise installed support, not a guarantee that a file can be read."""
    pdf = "pymupdf" if _available("pymupdf") else "pypdf" if _available("pypdf") else None
    return {
        "standard_library": ["txt", "md", "tex", "json", "html", "htm", "docx", "csv", "tsv"],
        "optional": {"pdf": {"available": pdf is not None, "backend": pdf},
                     "xlsx": {"available": _available("openpyxl"), "backend": "openpyxl"}},
        "ocr": False, "network": False, "executes_document_content": False,
        "limits": {"file_bytes": MAX_FILE_BYTES, "zip_uncompressed_bytes": MAX_ZIP_BYTES,
                   "zip_member_bytes": MAX_MEMBER_BYTES, "zip_members": MAX_ZIP_MEMBERS,
                   "output_characters": MAX_OUTPUT_CHARS, "blocks": MAX_BLOCKS,
                   "pdf_pages": MAX_PAGES, "sheet_rows": MAX_SHEET_ROWS,
                   "sheet_columns": MAX_SHEET_COLUMNS, "spreadsheet_cells": MAX_CELLS},
    }


class _Result:
    def __init__(self, fmt: str):
        self.data = {"format": fmt, "extraction_status": "complete", "blocks": [], "warnings": []}
        self.chars = 0
        self.full = False

    def warn(self, message: str, partial: bool = False):
        if message not in self.data["warnings"]:
            self.data["warnings"].append(message)
        if partial and self.data["extraction_status"] == "complete":
            self.data["extraction_status"] = "partial"

    def add(self, locator: str, value: str, *, preserve_empty=False):
        if self.full or (not value and not preserve_empty):
            return
        remaining = MAX_OUTPUT_CHARS - self.chars
        if len(self.data["blocks"]) >= MAX_BLOCKS or remaining <= 0:
            self.full = True
            self.warn("Output limit reached; later content was not extracted.", True)
            return
        if len(value) > remaining:
            value = value[:remaining]
            self.full = True
            self.warn("Output character limit reached; the last block is truncated and later content omitted.", True)
        self.data["blocks"].append({"locator": locator, "text": value})
        self.chars += len(value)

    def unsupported(self, message: str):
        self.data["extraction_status"] = "unsupported"
        self.warn(message)


def _text(path: Path, result: _Result):
    # No encoding guess: undecodable bytes must not become apparently reliable text.
    source = path.read_bytes().decode("utf-8-sig")
    for line_no, line in enumerate(source.splitlines(), 1):
        result.add(f"line[{line_no}]", line)
        if result.full:
            break


class _HTML(HTMLParser):
    IGNORE = {"script", "style", "template", "noscript"}
    VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self, result: _Result):
        super().__init__(convert_charrefs=True)
        self.result = result
        self.stack = []
        self.ignored = []
        self.ordinals = {}

    def handle_starttag(self, tag, attrs):
        if self.ignored:
            if tag not in self.VOID:
                self.ignored.append(tag)
            return
        if tag in self.IGNORE:
            self.ignored.append(tag)
            return
        if tag in {"img", "svg", "canvas", "audio", "video", "iframe", "object"}:
            self.result.warn("Embedded HTML media and external frames are not interpreted or fetched.", True)
        if tag not in self.VOID:
            self.stack.append(tag)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self.VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.ignored:
            if tag in self.ignored:
                index = len(self.ignored) - 1 - self.ignored[::-1].index(tag)
                del self.ignored[index:]
            return
        if tag in self.stack:
            index = len(self.stack) - 1 - self.stack[::-1].index(tag)
            del self.stack[index:]

    def handle_data(self, data):
        if self.ignored or not data.strip():
            return
        line, column = self.getpos()
        key = (line, column)
        self.ordinals[key] = self.ordinals.get(key, 0) + 1
        self.result.add(f"html-source:line[{line}]:column[{column + 1}]:text[{self.ordinals[key]}]", data.strip())


def _html(path: Path, result: _Result):
    result.warn("HTML locators refer to the saved source, not a rendered page. CSS visibility, layout and dynamic content are not evaluated.")
    parser = _HTML(result)
    with path.open("r", encoding="utf-8-sig") as source:
        while not result.full:
            chunk = source.read(65536)
            if not chunk:
                break
            parser.feed(chunk)
    parser.close()


def _safe_zip(path: Path) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(path)
    try:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_MEMBERS:
            raise ReaderLimit("Archive member count exceeds the safety limit.")
        names = set()
        total = 0
        for info in infos:
            if info.filename in names:
                raise ReaderLimit("Archive contains duplicate member names; ambiguous content refused.")
            names.add(info.filename)
            if info.flag_bits & 1:
                raise ReaderLimit("Encrypted ZIP members are unsupported.")
            if info.file_size > MAX_MEMBER_BYTES:
                raise ReaderLimit("Archive member exceeds the uncompressed size limit.")
            total += info.file_size
            if total > MAX_ZIP_BYTES:
                raise ReaderLimit("Archive exceeds the total uncompressed size limit.")
            if info.file_size > 1024 * 1024 and info.file_size / max(info.compress_size, 1) > 200:
                raise ReaderLimit("Archive member has an unsafe compression ratio.")
            # Office parsers may read extra parts. Check XML declarations everywhere.
            if info.filename.lower().endswith((".xml", ".rels")):
                _xml_bytes(archive.read(info))
        return archive
    except Exception:
        archive.close()
        raise


def _xml_bytes(content: bytes) -> bytes:
    if len(content) > MAX_MEMBER_BYTES:
        raise ReaderLimit("XML part exceeds the size limit.")
    # Strip NULs to also catch UTF-16/32 declarations before a parser expands entities.
    lowered = content.replace(b"\x00", b"").lower()
    if b"<!doctype" in lowered or b"<!entity" in lowered:
        raise ReaderLimit("XML DTD/entity declarations are not supported.")
    return content


def _word_text(element, result):
    pieces = []

    def walk(node):
        if node.tag in {W + "drawing", W + "pict", W + "txbxContent", W + "object"}:
            result.warn("DOCX drawings, text boxes and embedded objects are omitted.", True)
            return
        if node.tag in {W + "del", W + "moveFrom"}:
            result.warn("DOCX tracked revisions exist: deletions/move-from text omitted and insertions included; this is not a verified revision view.", True)
            return
        if node.tag in {W + "ins", W + "moveTo"}:
            result.warn("DOCX tracked revisions exist: deletions/move-from text omitted and insertions included; this is not a verified revision view.", True)
        if node.tag in {W + "fldChar", W + "instrText", W + "fldSimple"}:
            result.warn("DOCX fields are not evaluated; only any stored display text is included.", True)
        if node.tag == W + "t":
            pieces.append(node.text or "")
        elif node.tag == W + "tab":
            pieces.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            pieces.append("\n")
        elif node.tag in {W + "footnoteReference", W + "endnoteReference"}:
            kind = "footnote" if node.tag == W + "footnoteReference" else "endnote"
            pieces.append(f"[{kind}:{node.get(W + 'id', '?')}]")
        for child in node:
            walk(child)

    walk(element)
    return "".join(pieces)


def _word_container(container, base, result):
    counts = {}
    for child in container:
        local = child.tag.rsplit("}", 1)[-1]
        counts[local] = counts.get(local, 0) + 1
        locator = f"{base}/{local}[{counts[local]}]"
        if child.tag == W + "p":
            result.add(locator, _word_text(child, result), preserve_empty=True)
        elif child.tag in {W + "tbl", W + "tr", W + "tc", W + "sdt", W + "sdtContent", W + "customXml", W + "ins", W + "moveTo"}:
            _word_container(child, locator, result)
        if result.full:
            break


def _docx(path: Path, result: _Result):
    result.warn("DOCX extraction preserves selected text and XML element locators, not page numbers, typography, layout or full original fidelity. p/tbl/tr/tc mean paragraph/table/row/cell; indices are one-based among sibling elements of that type.")
    with _safe_zip(path) as archive:
        names = archive.namelist()
        if "word/document.xml" not in names:
            result.unsupported("Not a supported DOCX package: word/document.xml is missing.")
            return
        parts = ["word/document.xml"]
        parts += sorted(name for name in names if name.startswith(("word/header", "word/footer")) and name.endswith(".xml"))
        parts += [name for name in ("word/footnotes.xml", "word/endnotes.xml") if name in names]
        if any(name.startswith("word/comments") for name in names):
            result.warn("DOCX comments are not extracted.", True)
        if any(name.startswith(("word/media/", "word/embeddings/")) for name in names):
            result.warn("DOCX embedded images and objects are not interpreted.", True)
        for part in parts:
            root = ET.fromstring(_xml_bytes(archive.read(part)))
            if any(node.tag in {W + "ins", W + "del", W + "moveFrom", W + "moveTo", W + "pPrChange", W + "rPrChange", W + "tblPrChange"} for node in root.iter()):
                result.warn("DOCX contains tracked revisions; text projection does not resolve author acceptance or rejection.", True)
            if any(node.tag in {W + "altChunk", W + "subDoc"} for node in root.iter()):
                result.warn("DOCX linked/subdocument content is not extracted or fetched.", True)
            if any(node.tag in {W + "commentReference", W + "commentRangeStart", W + "commentRangeEnd"} for node in root.iter()):
                result.warn("DOCX comments are not extracted.", True)
            if any(node.tag == W + "sym" or node.tag.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/math}") for node in root.iter()):
                result.warn("DOCX mathematical markup and symbol-font glyphs are not interpreted.", True)
            if any(node.tag == W + "numPr" for node in root.iter()):
                result.warn("DOCX automatic list numbering is not expanded; consult the original for displayed list labels.", True)
            if any(node.tag == W + "hyperlink" for node in root.iter()):
                result.warn("DOCX hyperlink display text is included, but relationship target URLs are not extracted.", True)
            if part == "word/document.xml":
                body = root.find(W + "body")
                if body is None:
                    result.warn("DOCX document body is missing.", True)
                else:
                    _word_container(body, part + "/body", result)
            elif part.endswith(("footnotes.xml", "endnotes.xml")):
                for note in root:
                    note_id = note.get(W + "id", "unknown")
                    if note.get(W + "type") in {"separator", "continuationSeparator", "continuationNotice"} or note_id.startswith("-"):
                        continue
                    kind = note.tag.rsplit("}", 1)[-1]
                    _word_container(note, f"{part}/{kind}[id={note_id}]", result)
            else:
                _word_container(root, part, result)
            if result.full:
                break


def _csv(path: Path, result: _Result):
    result.warn("CSV/TSV locators use logical record and column numbers; records can span multiple physical lines. Formulas are text and are never evaluated.")
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source, delimiter="\t" if path.suffix.lower() == ".tsv" else ",", strict=True)
        for row_no, row in enumerate(reader, 1):
            for col_no, value in enumerate(row, 1):
                result.add(f"record[{row_no}]/column[{col_no}]", value)
                if result.full:
                    return


def _xlsx(path: Path, result: _Result):
    if not _available("openpyxl"):
        result.unsupported("XLSX requires the optional openpyxl package; no package was installed.")
        return
    actual_max_row, actual_max_col = 0, 0
    with _safe_zip(path) as archive:
        names = archive.namelist()
        if any(name.startswith(("xl/drawings/", "xl/media/", "xl/charts/", "xl/comments", "xl/embeddings/")) for name in names):
            result.warn("XLSX drawings, charts, images, comments and embedded objects are not extracted.", True)
        if any(name.startswith("xl/externalLinks/") for name in names):
            result.warn("XLSX external workbook links are not followed or refreshed.", True)
        # Read real cell coordinates rather than trusting the optional dimension cache.
        for name in names:
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                continue
            root = ET.fromstring(_xml_bytes(archive.read(name)))
            ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            for fallback_row, row in enumerate(root.iter(ns + "row"), 1):
                row_no = int(row.get("r", fallback_row))
                actual_max_row = max(actual_max_row, row_no)
                for fallback_col, cell in enumerate(row.findall(ns + "c"), 1):
                    match = re.fullmatch(r"([A-Za-z]+)([0-9]+)", cell.get("r", ""))
                    col_no = fallback_col
                    if match:
                        col_no = 0
                        for letter in match[1].upper():
                            col_no = col_no * 26 + ord(letter) - 64
                            if col_no > MAX_SHEET_COLUMNS:
                                break
                        actual_max_row = max(actual_max_row, int(match[2]))
                    actual_max_col = max(actual_max_col, col_no)
    if actual_max_row > MAX_SHEET_ROWS or actual_max_col > MAX_SHEET_COLUMNS:
        result.warn("Spreadsheet dimension limit reached; cells outside supported row/column bounds are omitted.", True)
    from openpyxl import load_workbook
    result.warn("Spreadsheet extraction reads cells, including hidden sheets, without reproducing formatting. Formulas are preserved as text, never evaluated; cached numeric results are not inferred.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    scanned = 0
    try:
        for sheet in workbook.worksheets:
            # Ignore potentially false dimension metadata, scan within verified bounds.
            sheet.reset_dimensions()
            if not actual_max_row or not actual_max_col:
                continue
            for row in sheet.iter_rows(max_row=min(actual_max_row, MAX_SHEET_ROWS), max_col=min(actual_max_col, MAX_SHEET_COLUMNS)):
                scanned += len(row)
                if scanned > MAX_CELLS:
                    result.warn("Spreadsheet cell scan limit reached; later cells are omitted.", True)
                    return
                nonempty = [cell for cell in row if cell.value is not None]
                for cell in nonempty:
                    value = cell.value
                    text = value.isoformat() if hasattr(value, "isoformat") else str(value)
                    quoted_sheet = sheet.title.replace("'", "''")
                    result.add(f"'{quoted_sheet}'!{cell.coordinate}", text)
                    if result.full:
                        return
    finally:
        workbook.close()


def _pdf(path: Path, result: _Result):
    backend = capabilities()["optional"]["pdf"]["backend"]
    if backend is None:
        result.unsupported("PDF requires optional pymupdf or pypdf; no package was installed. OCR is not provided.")
        return
    result.warn("PDF locators are one-based physical file pages, not printed page labels. Text extraction can misorder columns or tables; images, diagrams and visual layout are not interpreted. OCR is not provided.")
    empty_pages = []
    if backend == "pymupdf":
        import pymupdf
        with pymupdf.open(path) as document:
            if document.needs_pass:
                result.unsupported("Encrypted PDF requires a password and is unsupported.")
                return
            if len(document) > MAX_PAGES:
                result.warn("PDF page limit reached; later pages are omitted.", True)
            for index in range(min(len(document), MAX_PAGES)):
                page = document[index]
                content = page.get_text("text")
                if page.get_images():
                    result.warn("PDF contains page images that are not interpreted.", True)
                if not content.strip():
                    empty_pages.append(index + 1)
                result.add(f"physical-page[{index + 1}]", content)
                if result.full:
                    break
    else:
        from pypdf import PdfReader
        document = PdfReader(path, strict=True)
        if document.is_encrypted:
            result.unsupported("Encrypted PDF requires a password and is unsupported.")
            return
        if len(document.pages) > MAX_PAGES:
            result.warn("PDF page limit reached; later pages are omitted.", True)
        for index, page in enumerate(document.pages):
            if index >= MAX_PAGES:
                break
            content = page.extract_text() or ""
            if not content.strip():
                empty_pages.append(index + 1)
            result.add(f"physical-page[{index + 1}]", content)
            if result.full:
                break
    if empty_pages:
        preview = ", ".join(map(str, empty_pages[:20]))
        result.warn(f"No extractable text on {len(empty_pages)} PDF page(s), beginning with: {preview}. These pages may be scans or blank; no OCR was performed.", True)


def extract(path: Path) -> dict:
    """Extract without writes/network. Failures return unsupported or partial, never fabricated text."""
    path = Path(path)
    fmt = path.suffix.lower().lstrip(".") or "unknown"
    result = _Result(fmt)
    readers = {"txt": _text, "md": _text, "tex": _text, "json": _text,
               "html": _html, "htm": _html, "docx": _docx,
               "csv": _csv, "tsv": _csv, "xlsx": _xlsx, "pdf": _pdf}
    try:
        if not path.is_file():
            result.unsupported("Input is missing or is not a regular file.")
        elif path.stat().st_size > MAX_FILE_BYTES:
            result.unsupported("Input exceeds the file size safety limit.")
        elif fmt not in readers:
            result.unsupported("No reader for this format. The original may still be registered by Sources.")
        else:
            readers[fmt](path, result)
            if result.data["extraction_status"] == "complete" and not result.data["blocks"]:
                result.warn("No nonempty text was extracted; this does not establish that the original is empty.", True)
    except (OSError, ValueError, KeyError, RuntimeError, ImportError, zipfile.BadZipFile, ET.ParseError, csv.Error) as error:
        # Libraries expose heterogeneous parse errors; keep already extracted text explicitly partial.
        result.data["extraction_status"] = "partial" if result.data["blocks"] else "unsupported"
        result.warn(f"Reading failed ({type(error).__name__}): {str(error)[:300]}")
    except Exception as error:
        result.data["extraction_status"] = "partial" if result.data["blocks"] else "unsupported"
        result.warn(f"Optional reader failed ({type(error).__name__}): {str(error)[:300]}")
    return result.data
